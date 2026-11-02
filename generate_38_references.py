from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

# Author | article title | abbreviated journal | year | volume | pages/article number
DATA = """Geyer, R.; Jambeck, J. R.; Law, K. L.|Production, use, and fate of all plastics ever made.|Sci. Adv.|2017|3|e1700782
Mitrano, D. M.; Wick, P.; Nowack, B.|Placing nanoplastics in the context of global plastic pollution.|Nat. Nanotechnol.|2021|16|491-500
Vethaak, A. D.; Legler, J.|Microplastics and human health.|Science|2021|371|672-674
MacLeod, M.; Arp, H. P. H.; Tekman, M. B.; Jahnke, A.|The global threat from plastic pollution.|Science|2021|373|61-65
Jehanno, C.; Alty, J. W.; Roosen, M.; De Meester, S.; Dove, A. P.; Chen, E. Y. X.; Leibfarth, F. A.; Sardon, H.|Critical advances and future opportunities in upcycling commodity polymers.|Nature|2022|603|803-814
Zandieh, M.; Griffiths, E.; Waldie, A.; Li, S.; Honek, J.; Rezanezhad, F.; Van Cappellen, P.; Liu, J.|Catalytic and biocatalytic degradation of microplastics.|Exploration|2024|4|20230018
Ellis, L. D.; Rorrer, N. A.; Sullivan, K. P.; Otto, M.; McGeehan, J. E.; Román-Leshkov, Y.; Wierckx, N.; Beckham, G. T.|Chemical and biological catalysis for plastics recycling and upcycling.|Nat. Catal.|2021|4|539-556
Rahimi, A.; García, J. M.|Chemical recycling of waste plastics for new materials production.|Nat. Rev. Chem.|2017|1|0046
Zuin, V. G.; Kümmerer, K.|Chemistry and materials science for a sustainable circular polymeric economy.|Nat. Rev. Mater.|2022|7|76-78
Rybak, C.; Uyeda, C.|Advances in polymer chemistry enabled by transition metal catalysis.|Chem|2026|12|103140
Tournier, V.; Topham, C. M.; Gilles, A.; David, B.; Folgoas, C.; Moya-Leclair, E.; Kamionka, E.; Desrousseaux, M. L.; Texier, H.; Gavalda, S.; et al.|An engineered PET depolymerase to break down and recycle plastic bottles.|Nature|2020|580|216-219
Bell, E. L.; Smithson, R.; Kilbride, S.; Foster, J.; Hardy, F. J.; Ramachandran, S.; Tedstone, A. A.; Haigh, S. J.; Garforth, A. A.; Day, P. J. R.; et al.|Directed evolution of an efficient and thermostable PET depolymerase.|Nat. Catal.|2022|5|673-681
Wei, R.; Weber, G.; Blank, L. M.; Bornscheuer, U. T.|Process insights for harnessing biotechnology for plastic depolymerization.|Nat. Chem. Eng.|2025|2|110-117
Zimmermann, W.|Polyester-degrading enzymes in a circular economy of plastics.|Nat. Rev. Bioeng.|2025|3|681-696
Lu, H.; Diaz, D. J.; Czarnecki, N. J.; Zhu, C.; Kim, W.; Shroff, R.; Acosta, D. J.; Alexander, B. R.; Cole, H. O.; Zhang, Y.; et al.|Machine learning-aided engineering of hydrolases for PET depolymerization.|Nature|2022|604|662-667
Tournier, V.; Duquesne, S.; Guillamot, F. d. r.; Cramail, H.; Taton, D.; Marty, A.; André, I.|Enzymes’ Power for Plastics Degradation.|Chem. Rev.|2023|123|5612-5701
Bergeson, A. R.; Silvera, A. J.; Alper, H. S.|Bottlenecks in biobased approaches to plastic degradation.|Nat. Commun.|2024|15|4715
Rhys, G. G.|Artificial metalloenzymes in complex biological environments.|Nat. Chem. Biol.|2026|22|1565-1577
Zou, Z.; Higginson, B.; Ward, T. R.|Creation and optimization of artificial metalloenzymes: Harnessing the power of directed evolution and beyond.|Chem|2024|10|2373-2389
Vornholt, T.; Leiss-Maier, F.; Jeong, W. J.; Zeymer, C.; Song, W. J.; Roelfes, G.; Ward, T. R.|Artificial metalloenzymes.|Nat. Rev. Methods Primers|2024|4|78
Morita, I.; Ward, T. R.|Recent advances in the design and optimization of artificial metalloenzymes.|Curr. Opin. Chem. Biol.|2024|81|102508
Khatua, H.; Ghosh, A.; Das, S.; Patra, S.; Nandi, S.; Chattopadhyay, B.|Base Metal Catalysis in Nitrene Transfer Reactions.|Chem. Rev.|2026|126|1145-1205
Park, Y.; Kim, Y.; Chang, S.|Transition Metal-Catalyzed C–H Amination: Scope, Mechanism, and Applications.|Chem. Rev.|2017|117|9247-9301
Tran, B. L.; Li, B.; Driess, M.; Hartwig, J. F.|Copper-Catalyzed Intermolecular Amidation and Imidation of Unactivated Alkanes.|J. Am. Chem. Soc.|2014|136|2555-2563
Lee, J.; Jin, S.; Kim, D.; Hong, S. H.; Chang, S.|Cobalt-Catalyzed Intermolecular C–H Amidation of Unactivated Alkanes.|J. Am. Chem. Soc.|2021|143|5191-5200
Singh, R.; Mukherjee, A.|Metalloporphyrin Catalyzed C–H Amination.|ACS Catal.|2019|9|3604-3617
Chordia, S.; Narasimhan, S.; Lucini Paioni, A.; Baldus, M.; Roelfes, G.|In Vivo Assembly of Artificial Metalloenzymes and Application in Whole-Cell Biocatalysis.|Angew. Chem., Int. Ed.|2021|60|5913-5920
Yu, K.; Zou, Z.; Igareta, N. V.; Tachibana, R.; Bechter, J.; Köhler, V.; Chen, D.; Ward, T. R.|Artificial Metalloenzyme-Catalyzed Enantioselective Amidation via Nitrene Insertion in Unactivated C(sp3)–H Bonds.|J. Am. Chem. Soc.|2023|145|16621-16629
Dydio, P.; Key, H. M.; Hayashi, H.; Clark, D. S.; Hartwig, J. F.|Chemoselective, Enzymatic C–H Bond Amination Catalyzed by a Cytochrome P450 Containing an Ir(Me)-PIX Cofactor.|J. Am. Chem. Soc.|2017|139|1750-1753
Zawadiak, J.; Orlińska, B.; Marek, A. A.|Catalytic oxidation of polyethylene with oxygen in aqueous dispersion.|J. Appl. Polym. Sci.|2013|127|976-981
Zhao, H.; Zhang, D.; Loh, W. W.; Yam, S. J.; Pang, J. J. M.; Seng, D. H. L.; Png, Z. M.; Koh, X. Q.; Lee, Y. H.; Queh, L.; et al.|Aqueous-Based Biomimetic Porphyrin-Catalyzed Surface Oxidation of Polyethylenes for Enhanced 3D-Printability.|J. Am. Chem. Soc.|2026|148|22726-22738
Estell, D. A.; Graycar, T. P.; Miller, J. V.; Powers, D. B.; Wells, J. A.; Burnier, J. P.; Ng, P. G.|Probing Steric and Hydrophobic Effects on Enzyme-Substrate Interactions by Protein Engineering.|Science|1986|233|659-663
Austin, H. P.; Allen, M. D.; Donohoe, B. S.; Rorrer, N. A.; Kearns, F. L.; Silveira, R. L.; Pollard, B. C.; Dominick, G.; Duman, R.; El Omari, K.; et al.|Characterization and engineering of a plastic-degrading aromatic polyesterase.|Proc. Natl. Acad. Sci. U.S.A.|2018|115|E4350-E4357
Fischer, S.; Natter Perdiguero, A.; Lau, K.; Deliz Liang, A.|Hydrophobic tuning with non-canonical amino acids in a copper metalloenzyme.|Nat. Chem.|2026|18|1269-1277
Küchler, A.; Yoshimoto, M.; Luginbühl, S.; Mavelli, F.; Walde, P.|Enzymatic reactions in confined environments.|Nat. Nanotechnol.|2016|11|409-420
DelRe, C.; Jiang, Y.; Kang, P.; Kwon, J.; Hall, A.; Jayapurna, I.; Ruan, Z.; Ma, L.; Zolkin, K.; Li, T.; et al.|Near-complete depolymerization of polyesters with nano-dispersed enzymes.|Nature|2021|592|558-563
Woroch, C. P.; Klivansky, L.; McCarthy, D. P.; Miao, H.; Fink, Z.; Zuo, X.; Guthrie, C.; Wang, C.; Helms, B. A.; Hartwig, J. F.|Selective, Catalytic Oxidation of Polyethylene Surfaces.|Nature|2026||
Ciccia, N. R.; Shi, J. X.; Pal, S.; Hua, M.; Malollari, K. G.; Lizandara-Pueyo, C.; Risto, E.; Ernst, M.; Helms, B. A.; Messersmith, P. B.; et al.|Diverse functional polyethylenes by catalytic amination.|Science|2023|381|1433-1440"""

def create_docx(path):
    entries = [line.split('|') for line in DATA.splitlines()]
    assert len(entries) == 38 and all(len(row) == 6 for row in entries)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.left_margin = section.right_margin = Inches(1)
    style = doc.styles['Normal']
    style.font.name = 'Palatino Linotype'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)

    numbering = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), '50')
    mult = OxmlElement('w:multiLevelType')
    mult.set(qn('w:val'), 'singleLevel')
    abstract.append(mult)
    lvl = OxmlElement('w:lvl')
    lvl.set(qn('w:ilvl'), '0')
    for tag, val in [('start','1'), ('numFmt','decimal'), ('lvlText','%1.'), ('lvlJc','left')]:
        element = OxmlElement('w:'+tag)
        element.set(qn('w:val'), val)
        lvl.append(element)
    ppr = OxmlElement('w:pPr')
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'num')
    tab.set(qn('w:pos'), '540')
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '540')
    ind.set(qn('w:hanging'), '540')
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), '50')
    abstract_id = OxmlElement('w:abstractNumId')
    abstract_id.set(qn('w:val'), '50')
    num.append(abstract_id)
    numbering.append(num)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run('REFERENCES')
    heading_run.bold = True
    for authors, title, journal, year, volume, pages in entries:
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        num_pr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), '0')
        num_id = OxmlElement('w:numId')
        num_id.set(qn('w:val'), '50')
        num_pr.extend([ilvl, num_id])
        p._p.get_or_add_pPr().append(num_pr)
        p.add_run(authors+' ')
        for chunk in re.split(r'(?<![A-Za-z])(?=N[-‑–])', title):
            if chunk.startswith(('N-', 'N‑', 'N–')):
                p.add_run('N').italic = True
                p.add_run(chunk[1:])
            else:
                p.add_run(chunk)
        p.add_run(' ')
        p.add_run(journal).italic = True
        p.add_run(' ')
        p.add_run(year).bold = True
        if volume:
            p.add_run(', ')
            p.add_run(volume).italic = True
        if pages:
            p.add_run(', ' + re.sub(r'(?<=\w)-(?=\w)', '–', pages))
        p.add_run('.')
    doc.save(path)

if __name__ == '__main__':
    create_docx('ACS_References_38_Formatted.docx')
