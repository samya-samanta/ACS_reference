#!/usr/bin/env python3
"""Format numbered references in DOCX to the requested ACS presentation."""
import argparse
import re
from docx import Document
from docx.shared import Pt

JOURNALS = {
    "Journal of chemical education": "J. Chem. Educ.",
    "Chemical reviews": "Chem. Rev.",
    "Science": "Science",
    "Angewandte Chemie": "Angew. Chem.",
    "The Journal of Organic Chemistry": "J. Org. Chem.",
    "Berichte der deutschen chemischen Gesellschaft": "Ber. Dtsch. Chem. Ges.",
    "Journal of the American Chemical Society": "J. Am. Chem. Soc.",
    "ACS Catalysis": "ACS Catal.",
    "Organic Letters": "Org. Lett.",
    "Chemical Science": "Chem. Sci.",
    "Angewandte Chemie International Edition": "Angew. Chem., Int. Ed.",
    "Nature Communications": "Nat. Commun.",
    "Nature Chemistry": "Nat. Chem.",
    "Chemistry–A European Journal": "Chem. Eur. J.",
}
def parse_reference(ref):
    for original in sorted(JOURNALS, key=len, reverse=True):
        expression = r"(?<!\w)" + re.escape(original) + r"\s+(\d{4}),\s*(\d+),\s*(\d+\s*[-–]\s*\d+|e\w+|\d+)\.?\s*$"
        match = re.search(expression, ref, re.I)
        if match:
            return (ref[:match.start()], JOURNALS[original], *match.groups())
    return None

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("--output",default="ACS_formatted.docx")
    parser.add_argument("--limit",type=int,default=10)
    args=parser.parse_args()
    source=Document(args.input)
    dest=Document()
    dest.styles["Normal"].font.name="Times New Roman"
    dest.styles["Normal"].font.size=Pt(11)
    dest.add_heading("REFERENCES",1)
    count=0
    for item in source.paragraphs:
        entry=item.text.strip()
        if not re.match(r"^\d+\.\s",entry):
            continue
        if args.limit and count >= args.limit:
            break
        paragraph=dest.add_paragraph()
        paragraph.paragraph_format.left_indent=Pt(18)
        paragraph.paragraph_format.first_line_indent=Pt(-18)
        paragraph.paragraph_format.space_after=Pt(6)
        parts=parse_reference(entry)
        if parts:
            lead,journal,year,volume,pages=parts
            paragraph.add_run(lead)
            paragraph.add_run(journal).italic=True
            paragraph.add_run(" ")
            paragraph.add_run(year).bold=True
            paragraph.add_run(", ")
            paragraph.add_run(volume).italic=True
            paragraph.add_run(", "+re.sub(r"\s*[-–]\s*","–",pages)+".")
        else:
            paragraph.add_run(entry)
            print(f"REVIEW REQUIRED: reference {count+1} has an unknown journal or unexpected pattern")
        count+=1
    dest.save(args.output)
    print(f"Created {args.output} with {count} entries")

if __name__=="__main__":
    main()
